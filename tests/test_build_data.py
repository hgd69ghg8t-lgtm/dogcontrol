"""Tests for scripts/build_data.py using a small synthetic workbook laid out
like MCERT's "raw data" sheet. The numbers are made up, not MCERT data.

Run with: python -m unittest discover tests
"""

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import build_data  # noqa: E402

ROWS = [
    # council, category, subcategory, item, year, value
    ("Alpha District", "A_Registered dogs (total)", "all pure and cross breeds", "1001_total", 2024, 1000),
    ("Alpha District", "A_Registered dogs (total)", "all pure and cross breeds", "1001_total", 2025, 1100),
    ("Alpha District", "A_Registered dogs (total)", "all pure and cross breeds", "1006_classified as menacing", 2025, 11),
    ("New Zealand", "A_Registered dogs (total)", "all pure and cross breeds", "1001_total", 2025, 1100),
    ("Alpha District", "B_Registered pure breed", "aa_all (pure)", "1001_total", 2025, 400),
    ("Alpha District", "B_Registered pure breed", "Beagle", "1001_total", 2025, 30),
    ("Alpha District", "B_Registered pure breed", "Beagle", "1006_classified as menacing", 2025, 1),
    ("Alpha District", "C_Registered cross breed", "all (cross)", "1001_total", 2025, 700),
    ("Alpha District", "D_Owners", "total", "1010_currently registered dog owners", 2025, 800),
    # Mislabelled Act years in the source are summed into one metric.
    ("Alpha District", "E Ministry of Justice Statistics", "under the Dog Control Act 1996", "1012_total prosecutions", 2025, 2),
    ("Alpha District", "E Ministry of Justice Statistics", "under the Dog Control Act 1999", "1012_total prosecutions", 2025, 3),
    ("Alpha District", "F Accident Compensation Corporation Statistics", "Dog-related injury", "1014_number of new claims", 2024, "<4"),
    ("Alpha District", "F Accident Compensation Corporation Statistics", "Dog-related injury", "1014_number of new claims", 2025, 9),
    ("Alpha District", "F Accident Compensation Corporation Statistics", "Dog-related injury", "1016_total paid (excl. GST) for injuries this yr", 2025, 1234.567),
]


def make_workbook(path: Path) -> None:
    wb = Workbook()
    wb.active.title = "notes"
    ws = wb.create_sheet("raw data")
    ws.append(build_data.HEADER)
    for row in ROWS:
        ws.append(list(row))
    wb.save(path)


class BuildDataTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        path = Path(cls.tmp.name) / "stats.xlsx"
        make_workbook(path)
        cls.d = build_data.build(path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def value(self, metric, area, year):
        names = [a["name"] for a in self.d["areas"]]
        return self.d["values"][metric][names.index(area)][self.d["years"].index(year)]

    def test_areas_aggregates_first(self):
        self.assertEqual([(a["name"], a["type"]) for a in self.d["areas"]],
                         [("New Zealand", "national"), ("Alpha District", "council")])

    def test_metric_values(self):
        self.assertEqual(self.value("dogs", "Alpha District", 2025), 1100)
        self.assertEqual(self.value("dogs", "Alpha District", 2024), 1000)
        self.assertEqual(self.value("menacing", "Alpha District", 2025), 11)
        self.assertIsNone(self.value("menacing", "Alpha District", 2024))
        self.assertEqual(self.value("owners", "Alpha District", 2025), 800)
        self.assertEqual(self.value("acc_paid_year", "Alpha District", 2025), 1234.57)

    def test_breed_rows_do_not_leak_into_totals(self):
        self.assertEqual(self.value("dogs", "Alpha District", 2025), 1100)

    def test_mislabelled_act_rows_are_summed(self):
        self.assertEqual(self.value("prosecutions", "Alpha District", 2025), 5)

    def test_suppressed_values(self):
        self.assertIsNone(self.value("acc_claims", "Alpha District", 2024))
        self.assertEqual(self.value("acc_claims", "Alpha District", 2025), 9)
        self.assertIn(["acc_claims", 1, 2024, "<4"], self.d["suppressed"])

    def test_metric_year_ranges_and_unused_metrics_dropped(self):
        metrics = {m["id"]: m for m in self.d["metrics"]}
        self.assertEqual(metrics["dogs"]["years"], [2024, 2025])
        self.assertEqual(metrics["acc_claims"]["years"], [2025, 2025])
        self.assertNotIn("dangerous", metrics)

    def test_breeds(self):
        b = self.d["breeds"]
        self.assertEqual(b["names"], ["All pure breeds", "Beagle", "Cross breeds"])
        self.assertEqual(b["kinds"], ["all", "breed", "cross"])
        self.assertEqual(b["years"], [2025])
        beagle = b["names"].index("Beagle")
        self.assertEqual(b["values"]["dogs"][beagle][1][0], 30)
        self.assertEqual(b["values"]["menacing"][beagle][1][0], 1)
        self.assertIsNone(b["values"]["menacing"][beagle][0][0])

    def test_rejects_unexpected_layout(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bad.xlsx"
            wb = Workbook()
            wb.active.title = "raw data"
            wb.active.append(["Council", "Value"])
            wb.save(path)
            with self.assertRaises(SystemExit):
                build_data.build(path)


if __name__ == "__main__":
    unittest.main()
