#!/usr/bin/env python3
"""Convert the MCERT dog control statistics workbook(s) into JSON for the site.

Reads every .xlsx in data/raw/ and writes site/data/datasets.json. Each
worksheet becomes one dataset: a header row is detected automatically, the
first mostly-text column is used as the row label (e.g. territorial
authority), and every mostly-numeric column becomes a selectable measure.

Usage:
    python scripts/build_data.py [--raw data/raw] [--out site/data/datasets.json]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TOTAL_RE = re.compile(r"^\s*(grand\s+)?total\b|\btotals?\s*$", re.IGNORECASE)


def clean(value):
    """Normalise a cell value to str / int / float / None."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value, 4)
    if isinstance(value, int):
        return value
    text = re.sub(r"\s+", " ", str(value)).strip()
    if not text:
        return None
    numeric = text.replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", numeric):
        number = float(numeric)
        return int(number) if number.is_integer() else number
    return text


def is_number(value) -> bool:
    return isinstance(value, (int, float))


def find_header_row(rows: list[list]) -> int | None:
    """Index of the header row: a row of mostly text followed by data rows
    that contain at least one number."""
    for i, row in enumerate(rows[:40]):
        filled = [v for v in row if v is not None]
        if len(filled) < 2:
            continue
        if sum(isinstance(v, str) for v in filled) / len(filled) < 0.6:
            continue
        below = [r for r in rows[i + 1 : i + 6] if any(v is not None for v in r)]
        if below and any(any(is_number(v) for v in r) for r in below):
            return i
    return None


def unique_headers(raw_headers: list, width: int) -> list[str]:
    headers, seen = [], {}
    for idx in range(width):
        name = raw_headers[idx] if idx < len(raw_headers) else None
        name = str(name) if name is not None else f"Column {idx + 1}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        headers.append(name if count == 0 else f"{name} ({count + 1})")
    return headers


def sheet_to_dataset(ws, source: str) -> dict | None:
    rows = [[clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(v is not None for v in r)]
    if not rows:
        return None

    header_idx = find_header_row(rows)
    if header_idx is None:
        return None

    # Text above the header (title, notes) is kept as context for the page.
    notes = [
        " ".join(str(v) for v in r if v is not None) for r in rows[:header_idx]
    ]
    body = rows[header_idx + 1 :]
    width = max(len(r) for r in [rows[header_idx], *body])
    headers = unique_headers(rows[header_idx], width)
    body = [r + [None] * (width - len(r)) for r in body]

    # Drop columns that are completely empty.
    keep = [c for c in range(width) if any(r[c] is not None for r in body)]
    headers = [headers[c] for c in keep]
    body = [[r[c] for c in keep] for r in body]

    numeric, text = [], []
    for c, name in enumerate(headers):
        values = [r[c] for r in body if r[c] is not None]
        if values and sum(is_number(v) for v in values) / len(values) >= 0.8:
            numeric.append(name)
        else:
            text.append(name)
    if not numeric:
        return None

    label = text[0] if text else headers[0]
    label_idx = headers.index(label)

    records, footnotes = [], []
    for r in body:
        has_numbers = any(is_number(r[headers.index(n)]) for n in numeric)
        if not has_numbers:
            # Rows with text but no numbers after the table are usually notes.
            footnotes.append(" ".join(str(v) for v in r if v is not None))
            continue
        record = dict(zip(headers, r))
        record["_total"] = bool(
            isinstance(r[label_idx], str) and TOTAL_RE.search(r[label_idx])
        )
        records.append(record)

    return {
        "id": re.sub(r"[^a-z0-9]+", "-", f"{source}-{ws.title}".lower()).strip("-"),
        "source": source,
        "sheet": ws.title,
        "notes": [n for n in notes if n],
        "footnotes": [f for f in footnotes if f],
        "label": label,
        "dimensions": text,
        "measures": [m for m in numeric if m != label],
        "columns": headers,
        "rows": records,
    }


def build(raw_dir: Path) -> dict:
    datasets = []
    files = sorted(p for p in raw_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    for path in files:
        wb = load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            ds = sheet_to_dataset(ws, path.stem)
            if ds:
                datasets.append(ds)
        wb.close()
    return {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "files": [p.name for p in files],
        "datasets": datasets,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--raw", type=Path, default=ROOT / "data" / "raw")
    parser.add_argument("--out", type=Path, default=ROOT / "site" / "data" / "datasets.json")
    args = parser.parse_args(argv)

    result = build(args.raw)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=1, ensure_ascii=False) + "\n")

    if not result["files"]:
        print(f"No .xlsx files found in {args.raw}; wrote an empty dataset list.")
    for ds in result["datasets"]:
        print(f"{ds['source']} / {ds['sheet']}: {len(ds['rows'])} rows, "
              f"{len(ds['measures'])} measures (label: {ds['label']})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
